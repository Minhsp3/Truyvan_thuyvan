import io

import pandas as pd
from openpyxl import Workbook

from data_utils import (
    PLANTS,
    _historical_cell_map,
    _set_numeric_cells_in_worksheet,
    analog_similarity_series,
    annual_projection_table,
    comparison_table,
    current_month_actuals,
    current_period,
    latest_complete_actual_date,
    monthly_cumulative_production,
    prepare_analog_comparison_data,
    rank_analog_years,
    read_cumulative_hydrology_report,
    read_workbook,
)


def sample_data() -> pd.DataFrame:
    rows = []
    values = {
        2023: ((90, 180), (45, 90), (70, 140)),
        2024: ((101, 201), (51, 99), (81, 161)),
        2026: ((100, 200), (50, 100), (80, 160)),
    }
    plants = ["Buôn Kuốp", "Buôn Tua Srah", "Srêpốk 3"]
    for year, plant_values in values.items():
        for plant, (production, flow) in zip(plants, plant_values):
            rows.append(
                {
                    "year": year,
                    "month": 9,
                    "plant": plant,
                    "production": production,
                    "flow": flow,
                }
            )
    return pd.DataFrame(rows)


def test_rank_analog_years_selects_closest_year() -> None:
    ranking = rank_analog_years(sample_data(), 2026, 9)
    assert int(ranking.iloc[0]["year"]) == 2024
    assert ranking.iloc[0]["comparisons"] == 6


def test_plant_ranking_and_comparison() -> None:
    ranking = rank_analog_years(sample_data(), 2026, 9, "Buôn Kuốp")
    table = comparison_table(sample_data(), 2026, 9, int(ranking.iloc[0]["year"]), "Buôn Kuốp")
    assert int(ranking.iloc[0]["year"]) == 2024
    assert table["Nhà máy"].tolist() == ["Buôn Kuốp"]
    assert table["Tháng"].tolist() == [9]


def test_similarity_series_contains_both_years() -> None:
    chart = analog_similarity_series(sample_data(), 2026, 9, "Buôn Kuốp", 2024)
    assert chart["Năm"].tolist() == ["2026", "2024"]
    assert chart["Tháng"].tolist() == [9, 9]
    assert chart.loc[chart["Năm"] == "2026", "Chỉ số tương đồng"].iloc[0] == 100


def test_ranking_uses_all_months_from_january_to_cutoff() -> None:
    rows = []
    # 2023 matches the cutoff month exactly but is very different in January.
    # 2024 is close over both months and must therefore rank first.
    values = {
        2023: {1: (300, 600), 2: (110, 210)},
        2024: {1: (98, 198), 2: (108, 208)},
        2026: {1: (100, 200), 2: (110, 210)},
    }
    for year, monthly in values.items():
        for month, (production, flow) in monthly.items():
            rows.append(
                {
                    "year": year,
                    "month": month,
                    "plant": "Buôn Kuốp",
                    "production": production,
                    "flow": flow,
                }
            )

    ranking = rank_analog_years(pd.DataFrame(rows), 2026, 2, "Buôn Kuốp")
    assert int(ranking.iloc[0]["year"]) == 2024
    assert int(ranking.iloc[0]["comparisons"]) == 4


def test_current_period_falls_back_to_latest() -> None:
    year, month, exact = current_period(sample_data(), pd.Timestamp("2026-10-01"))
    assert (year, month, exact) == (2026, 9, False)


def test_partial_month_uses_same_period_daily_values_for_analog_ranking() -> None:
    data = pd.DataFrame(
        [
            {
                "year": 2026,
                "month": 10,
                "plant": "Buôn Kuốp",
                "production": 130.0,
                "flow": 200.0,
            },
            {
                "year": 2025,
                "month": 10,
                "plant": "Buôn Kuốp",
                "production": 300.0,
                "flow": 190.0,
            },
        ]
    )

    daily_production = pd.DataFrame(
        [
            {
                "date": pd.Timestamp(year=2025, month=10, day=day),
                "plant": "Buôn Kuốp",
                "production": float(day),
            }
            for day in range(1, 15)
        ]
    )
    daily_hydrology = pd.DataFrame(
        [
            {
                "date": pd.Timestamp(year=2025, month=10, day=day),
                "plant": "Buôn Kuốp",
                "flow": float(100 + day),
            }
            for day in range(1, 15)
        ]
    )

    comparison = prepare_analog_comparison_data(
        data,
        2026,
        10,
        pd.Timestamp("2026-10-13"),
        daily_production,
        daily_hydrology,
    )

    current = comparison[(comparison["year"] == 2026) & (comparison["month"] == 10)].iloc[0]
    historical = comparison[comparison["year"] == 2025].iloc[0]
    assert current["production"] == 130.0
    assert current["flow"] == 200.0
    assert historical["production"] == sum(range(1, 14))
    assert historical["flow"] == sum(range(101, 114)) / 13


def test_annual_projection_uses_actual_then_analog_values() -> None:
    rows = []
    for plant in PLANTS:
        for month in range(1, 13):
            rows.append(
                {
                    "year": 2024,
                    "month": month,
                    "plant": plant,
                    "production": month * 10,
                    "flow": month * 20,
                }
            )
        for month in (1, 2):
            rows.append(
                {
                    "year": 2026,
                    "month": month,
                    "plant": plant,
                    "production": 100 + month,
                    "flow": 200 + month,
                }
            )

    table = annual_projection_table(pd.DataFrame(rows), 2026, 2)
    first_plant = table[table["Nhà máy"] == PLANTS[0]].iloc[0]
    total = table[table["Nhà máy"] == "Tổng cộng"].iloc[0]

    assert int(first_plant["Năm quy luật"]) == 2024
    assert first_plant["T1"] == 101
    assert first_plant["T2"] == 102
    assert first_plant["T3"] == 30
    assert first_plant["Tổng năm"] == 953
    assert total["T1"] == 303
    assert total["Tổng năm"] == 2859


def test_partial_month_projection_adds_analog_remainder() -> None:
    rows = []
    for plant in PLANTS:
        for month in range(1, 13):
            rows.append(
                {
                    "year": 2024,
                    "month": month,
                    "plant": plant,
                    "production": month * 10,
                    "flow": month * 20,
                }
            )
        for month in range(1, 9):
            rows.append(
                {
                    "year": 2026,
                    "month": month,
                    "plant": plant,
                    "production": month * 10,
                    "flow": month * 20,
                }
            )
        rows.append(
            {
                "year": 2026,
                "month": 9,
                "plant": plant,
                "production": 63,
                "flow": float("nan"),
            }
        )

    table = annual_projection_table(
        pd.DataFrame(rows), 2026, 9, pd.Timestamp("2026-09-21")
    )
    first_plant = table[table["Nhà máy"] == PLANTS[0]].iloc[0]
    assert first_plant["T9"] == 90
    assert first_plant["T10"] == 100


def test_partial_month_projection_sums_each_remaining_analog_day() -> None:
    rows = []
    daily_rows = []
    for plant in PLANTS:
        for month in range(1, 13):
            rows.append(
                {
                    "year": 2024,
                    "month": month,
                    "plant": plant,
                    "production": month * 10,
                    "flow": month * 20,
                }
            )
        for month in range(1, 9):
            rows.append(
                {
                    "year": 2026,
                    "month": month,
                    "plant": plant,
                    "production": month * 10,
                    "flow": month * 20,
                }
            )
        rows.append(
            {"year": 2026, "month": 9, "plant": plant, "production": 63, "flow": 180}
        )
        for day in range(22, 31):
            daily_rows.append(
                {
                    "date": pd.Timestamp(year=2024, month=9, day=day),
                    "plant": plant,
                    "production": 1.0,
                }
            )

    table = annual_projection_table(
        pd.DataFrame(rows),
        2026,
        9,
        pd.Timestamp("2026-09-21"),
        pd.DataFrame(daily_rows),
    )
    first_plant = table[table["Nhà máy"] == PLANTS[0]].iloc[0]
    assert first_plant["T9"] == 72


def test_cumulative_report_reads_monthly_inflow_from_column_d() -> None:
    rows = [[None] * 5 for _ in range(16)]
    rows[0][0] = "Đến ngày 22/09/2026"
    expected = {"Buôn Kuốp": 258.38, "Buôn Tua Srah": 246.98, "Srêpốk 3": 317.34}
    for start_row, (plant, flow) in zip((2, 7, 12), expected.items()):
        rows[start_row][1] = plant
        rows[start_row + 1][1] = "Lưu lượng TB vào hồ"
        rows[start_row + 1][2] = "m3/s"
        rows[start_row + 1][3] = flow

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame(rows).to_excel(writer, sheet_name="Luy ke", index=False, header=False)
    buffer.seek(0)

    parsed = read_workbook(buffer)
    september = parsed[(parsed["year"] == 2026) & (parsed["month"] == 9)]
    actual = september.set_index("plant")["flow"].to_dict()
    assert actual == expected

    buffer.seek(0)
    dated = read_cumulative_hydrology_report(buffer)
    assert dated["date"].nunique() == 1
    assert dated["date"].iloc[0] == pd.Timestamp("2026-09-22")


def test_current_actuals_use_dated_cumulative_report() -> None:
    production = pd.DataFrame(
        [
            {
                "date": pd.Timestamp("2026-10-13"),
                "plant": plant,
                "production": float(index + 1),
            }
            for index, plant in enumerate(PLANTS)
        ]
    )
    report = pd.DataFrame(
        [
            {"date": pd.Timestamp("2026-10-13"), "plant": plant, "flow": 100 + index}
            for index, plant in enumerate(PLANTS)
        ]
    )
    empty_hydrology = pd.DataFrame(columns=["date", "plant", "flow"])

    as_of = latest_complete_actual_date(production, empty_hydrology, report)
    actuals = current_month_actuals(production, empty_hydrology, report, as_of)

    assert as_of == pd.Timestamp("2026-10-13")
    assert actuals["production"].tolist() == [1.0, 2.0, 3.0]
    assert actuals["flow"].tolist() == [100.0, 101.0, 102.0]


def test_monthly_cumulative_production_compares_only_through_same_day() -> None:
    daily = pd.DataFrame(
        [
            {
                "date": pd.Timestamp(year=2025, month=10, day=day),
                "plant": plant,
                "production": float(day + plant_index),
            }
            for plant_index, plant in enumerate(PLANTS)
            for day in range(1, 15)
        ]
    )

    result = monthly_cumulative_production(daily, 2025, 10, 13).set_index("plant")

    assert result.at["Buôn Kuốp", "production"] == sum(range(1, 14))
    assert result.at["Buôn Tua Srah", "production"] == sum(range(2, 15))
    assert result.at["Srêpốk 3", "production"] == sum(range(3, 16))
    assert result["days"].tolist() == [13, 13, 13]


def test_dual_month_block_workbook_is_parsed() -> None:
    columns = 28
    rows = [[None] * columns for _ in range(4)]
    rows[0][2] = "Sản lượng [tr.kWh]"
    rows[0][16] = "Lưu lượng [m3/s]"
    for month in range(1, 13):
        rows[1][month + 1] = month
        rows[1][month + 15] = month
        rows[2][month + 1] = month * 10
        rows[2][month + 15] = month * 20
    rows[2][1] = 2026
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame(rows).to_excel(writer, sheet_name="BK", index=False, header=False)
    buffer.seek(0)

    parsed = read_workbook(buffer)
    september = parsed[(parsed["year"] == 2026) & (parsed["month"] == 9)].iloc[0]
    assert september["plant"] == "Buôn Kuốp"
    assert september["production"] == 90
    assert september["flow"] == 180


def test_historical_cell_map_finds_year_and_both_month_blocks() -> None:
    workbook = Workbook()
    sheet = workbook.active
    for month in range(1, 13):
        sheet.cell(2, month + 2).value = month
        sheet.cell(2, month + 16).value = month
    sheet.cell(3, 2).value = 2026

    row, production_columns, flow_columns = _historical_cell_map(sheet, 2026)

    assert row == 3
    assert production_columns[10] == 12
    assert flow_columns[10] == 26


def test_xlsx_cell_patch_preserves_other_formula_cache() -> None:
    worksheet_xml = b"""<?xml version="1.0" encoding="UTF-8"?>
    <worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
      <sheetData><row r="3">
        <c r="C3"><f>Q3*2</f><v>54.105</v></c>
        <c r="L52" t="n"><v>83.582</v></c>
        <c r="Z52" t="n"><v>332.23</v></c>
      </row></sheetData>
    </worksheet>"""

    patched = _set_numeric_cells_in_worksheet(
        worksheet_xml, {"L52": 83.5827994, "Z52": 332.23}
    )

    assert b"<f>Q3*2</f><v>54.105</v>" in patched
    assert b'<c r="L52" t="n"><v>83.5827994</v></c>' in patched
