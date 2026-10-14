from __future__ import annotations

import io
import os
import posixpath
import re
import shutil
import unicodedata
import uuid
import zipfile
from datetime import datetime
from pathlib import Path
from typing import BinaryIO, Iterable
from xml.etree import ElementTree as ET

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


PLANTS = ["Buôn Kuốp", "Buôn Tua Srah", "Srêpốk 3"]

HISTORICAL_WORKBOOKS = {
    "Buôn Kuốp": "Lưu lượng Sản lượng Buôn Kuop 1977_2026.xlsx",
    "Buôn Tua Srah": "Lưu lượng Sản lượng Buôn Tua Srah 1977_2026.xlsx",
    "Srêpốk 3": "Lưu lượng Sản lượng Srepok3 1977_2026.xlsx",
}


def _plain(value: object) -> str:
    text = unicodedata.normalize("NFD", str(value).strip().lower())
    return "".join(char for char in text if unicodedata.category(char) != "Mn").replace("đ", "d")


def canonical_plant(value: object) -> str | None:
    text = re.sub(r"[^a-z0-9]", "", _plain(value))
    aliases = {
        "Buôn Kuốp": ("buonkuop", "buonkuop", "bk"),
        "Buôn Tua Srah": ("buontuasrah", "bts"),
        "Srêpốk 3": ("srepok3", "srepok3", "sp3"),
    }
    for plant, keys in aliases.items():
        if any(text == key or key in text for key in keys):
            return plant
    return None


def _number(value: object) -> float:
    if pd.isna(value):
        return np.nan
    if isinstance(value, str):
        value = value.strip().replace(" ", "").replace(",", ".")
    try:
        return float(value)
    except (TypeError, ValueError):
        return np.nan


def _date_series(values: pd.Series) -> pd.Series:
    """Parse ISO dates without swapping month/day and Vietnamese dates day-first."""
    result = pd.Series(pd.NaT, index=values.index, dtype="datetime64[ns]")
    text = values.astype(str).str.strip()
    iso_mask = text.str.fullmatch(r"\d{4}-\d{1,2}-\d{1,2}")
    result.loc[iso_mask] = pd.to_datetime(
        text.loc[iso_mask], format="%Y-%m-%d", errors="coerce"
    )
    result.loc[~iso_mask] = pd.to_datetime(
        values.loc[~iso_mask], errors="coerce", dayfirst=True
    )
    return result


def _parse_year_sheets(source: str | Path | BinaryIO) -> pd.DataFrame:
    """Parse workbooks whose sheets are years and columns are months 1..12."""
    excel = pd.ExcelFile(source)
    records: list[dict[str, object]] = []
    for sheet in excel.sheet_names:
        match = re.fullmatch(r"\s*((?:19|20)\d{2})\s*", str(sheet))
        if not match:
            continue
        year = int(match.group(1))
        raw = pd.read_excel(excel, sheet_name=sheet, header=None)
        if raw.empty or raw.shape[1] < 3:
            continue

        header_row = None
        month_columns: dict[int, int] = {}
        for row_idx in range(min(12, len(raw))):
            candidate: dict[int, int] = {}
            for col_idx, value in enumerate(raw.iloc[row_idx]):
                number = _number(value)
                if np.isfinite(number) and number.is_integer() and 1 <= number <= 12:
                    candidate[int(number)] = col_idx
                else:
                    text = _plain(value)
                    month_match = re.search(r"thang\s*(1[0-2]|[1-9])", text)
                    if month_match:
                        candidate[int(month_match.group(1))] = col_idx
            if len(candidate) >= 6:
                header_row, month_columns = row_idx, candidate
                break
        if header_row is None:
            continue

        current_metric: str | None = None
        for row_idx in range(header_row + 1, len(raw)):
            row = raw.iloc[row_idx]
            label_text = " ".join(_plain(value) for value in row.iloc[:3] if not pd.isna(value))
            if "san luong" in label_text or "tr.kwh" in label_text or "mwh" in label_text:
                current_metric = "production"
            elif "qv" in label_text or "luu luong" in label_text or "m3/s" in label_text:
                current_metric = "flow"

            plant = next((canonical_plant(value) for value in row.iloc[:4] if canonical_plant(value)), None)
            if not plant or not current_metric:
                continue
            for month, col_idx in month_columns.items():
                value = _number(row.iloc[col_idx]) if col_idx < len(row) else np.nan
                if np.isfinite(value):
                    records.append(
                        {"year": year, "month": month, "plant": plant, current_metric: value}
                    )
    return _combine_metric_rows(pd.DataFrame(records))


def _find_column(columns: Iterable[object], keywords: tuple[str, ...]) -> object | None:
    for column in columns:
        name = _plain(column)
        if any(keyword in name for keyword in keywords):
            return column
    return None


def _parse_long_form(source: str | Path | BinaryIO) -> pd.DataFrame:
    """Parse a conventional table with plant/date or plant/year/month columns."""
    excel = pd.ExcelFile(source)
    frames: list[pd.DataFrame] = []
    for sheet in excel.sheet_names:
        for header in range(0, 8):
            try:
                df = pd.read_excel(excel, sheet_name=sheet, header=header)
            except Exception:
                continue
            plant_col = _find_column(df.columns, ("nha may", "plant", "cong trinh"))
            prod_col = _find_column(df.columns, ("san luong", "production", "dien luong"))
            flow_col = _find_column(df.columns, ("luu luong", "qve", "q ve", "qv", "flow"))
            if plant_col is None or (prod_col is None and flow_col is None):
                continue
            date_col = _find_column(df.columns, ("ngay", "date", "thoi gian"))
            year_col = _find_column(df.columns, ("nam", "year"))
            month_col = _find_column(df.columns, ("thang", "month"))
            if date_col is None and (year_col is None or month_col is None):
                continue

            out = pd.DataFrame({"plant": df[plant_col].map(canonical_plant)})
            if date_col is not None:
                dates = pd.to_datetime(df[date_col], errors="coerce", dayfirst=True)
                out["year"], out["month"] = dates.dt.year, dates.dt.month
            else:
                out["year"] = pd.to_numeric(df[year_col], errors="coerce")
                out["month"] = pd.to_numeric(df[month_col], errors="coerce")
            if prod_col is not None:
                out["production"] = pd.to_numeric(df[prod_col], errors="coerce")
            if flow_col is not None:
                out["flow"] = pd.to_numeric(df[flow_col], errors="coerce")
            frames.append(out)
            break
    return _combine_metric_rows(pd.concat(frames, ignore_index=True)) if frames else pd.DataFrame()


def _parse_year_rows(source: str | Path | BinaryIO) -> pd.DataFrame:
    """Parse sheets with one year per row and month columns, commonly used for flow."""
    excel = pd.ExcelFile(source)
    records: list[dict[str, object]] = []
    for sheet in excel.sheet_names:
        plant = canonical_plant(sheet)
        raw = pd.read_excel(excel, sheet_name=sheet, header=None)
        if plant is None:
            for value in raw.iloc[:5, :5].to_numpy().ravel():
                plant = canonical_plant(value)
                if plant:
                    break
        if not plant:
            continue
        for row_idx in range(min(12, len(raw))):
            labels = [_plain(value) for value in raw.iloc[row_idx]]
            if not any(label == "nam" for label in labels):
                continue
            year_col = next(i for i, label in enumerate(labels) if label == "nam")
            month_columns: dict[int, int] = {}
            for col_idx, label in enumerate(labels):
                match = re.search(r"thang\s*(1[0-2]|[1-9])", label)
                if match:
                    month_columns[int(match.group(1))] = col_idx
            if len(month_columns) < 6:
                continue
            for data_idx in range(row_idx + 1, len(raw)):
                year = _number(raw.iloc[data_idx, year_col])
                if not np.isfinite(year) or not 1900 <= year <= 2100:
                    continue
                for month, col_idx in month_columns.items():
                    flow = _number(raw.iloc[data_idx, col_idx])
                    if np.isfinite(flow):
                        records.append(
                            {"year": int(year), "month": month, "plant": plant, "flow": flow}
                        )
            break
    return _combine_metric_rows(pd.DataFrame(records))


def _parse_dual_matrix(source: str | Path | BinaryIO) -> pd.DataFrame:
    """Parse one-plant sheets with production and flow as two 12-month blocks."""
    excel = pd.ExcelFile(source)
    records: list[dict[str, object]] = []
    for sheet in excel.sheet_names:
        plant = canonical_plant(sheet)
        if not plant:
            continue
        raw = pd.read_excel(excel, sheet_name=sheet, header=None)
        if raw.empty:
            continue
        for header_idx in range(min(10, len(raw))):
            month_positions: dict[int, list[int]] = {month: [] for month in range(1, 13)}
            for col_idx, value in enumerate(raw.iloc[header_idx]):
                number = _number(value)
                if np.isfinite(number) and number.is_integer() and 1 <= number <= 12:
                    month_positions[int(number)].append(col_idx)
            if not all(len(month_positions[month]) >= 2 for month in range(1, 13)):
                continue
            production_cols = {month: month_positions[month][0] for month in range(1, 13)}
            flow_cols = {month: month_positions[month][1] for month in range(1, 13)}
            first_month_col = production_cols[1]
            year_col = first_month_col - 1
            if year_col < 0:
                continue
            for row_idx in range(header_idx + 1, len(raw)):
                year = _number(raw.iloc[row_idx, year_col])
                if not np.isfinite(year) or not year.is_integer() or not 1900 <= year <= 2100:
                    continue
                for month in range(1, 13):
                    production = _number(raw.iloc[row_idx, production_cols[month]])
                    flow = _number(raw.iloc[row_idx, flow_cols[month]])
                    if np.isfinite(production) or np.isfinite(flow):
                        records.append(
                            {
                                "year": int(year),
                                "month": month,
                                "plant": plant,
                                "production": production,
                                "flow": flow,
                            }
                        )
            break
    return _combine_metric_rows(pd.DataFrame(records))


def _parse_daily_production_report(
    source: str | Path | BinaryIO, cutoff_date: pd.Timestamp | None = None
) -> pd.DataFrame:
    """Parse the daily BCSX report and aggregate generator output to million kWh."""
    daily = read_daily_production_workbook(source, cutoff_date)
    if daily.empty:
        return _combine_metric_rows(pd.DataFrame())
    monthly = daily.assign(year=daily["date"].dt.year, month=daily["date"].dt.month)
    monthly = monthly.groupby(["year", "month", "plant"], as_index=False)["production"].sum()
    return _combine_metric_rows(monthly)


def read_daily_production_workbook(
    source: str | Path | BinaryIO, cutoff_date: pd.Timestamp | None = None
) -> pd.DataFrame:
    """Read daily generator output and convert MWh to million kWh."""
    excel = pd.ExcelFile(source)
    records: list[dict[str, object]] = []
    cutoff = pd.Timestamp(cutoff_date).normalize() if cutoff_date is not None else None
    for sheet in excel.sheet_names:
        raw = pd.read_excel(excel, sheet_name=sheet, header=None)
        if raw.empty:
            continue
        header_row = next(
            (
                row_idx
                for row_idx in range(min(12, len(raw)))
                if any(_plain(value) == "ngay" for value in raw.iloc[row_idx] if pd.notna(value))
            ),
            None,
        )
        if header_row is None or header_row + 1 >= len(raw):
            continue

        group_headers = raw.iloc[header_row].ffill()
        metric_headers = raw.iloc[header_row + 1]
        production_columns: dict[str, int] = {}
        for col_idx in range(1, raw.shape[1]):
            plant = canonical_plant(group_headers.iloc[col_idx])
            metric = _plain(metric_headers.iloc[col_idx])
            if plant and "dau cuc" in metric and "mwh" in metric:
                production_columns[plant] = col_idx
        if not production_columns:
            continue

        dates = _date_series(raw.iloc[header_row + 2 :, 0])
        valid_dates = dates.notna()
        if cutoff is not None:
            valid_dates &= dates.dt.normalize().le(cutoff)
        for plant, col_idx in production_columns.items():
            values = pd.to_numeric(raw.iloc[header_row + 2 :, col_idx], errors="coerce")
            daily = pd.DataFrame({"date": dates, "production_mwh": values})[valid_dates]
            daily = daily.dropna(subset=["production_mwh"])
            for _, daily_row in daily.iterrows():
                records.append(
                    {
                        "date": pd.Timestamp(daily_row["date"]).normalize(),
                        "plant": plant,
                        "production": float(daily_row["production_mwh"]) / 1000.0,
                    }
                )
    if not records:
        return pd.DataFrame(columns=["date", "plant", "production"])
    return (
        pd.DataFrame(records)
        .groupby(["date", "plant"], as_index=False)["production"]
        .mean()
        .sort_values(["date", "plant"])
        .reset_index(drop=True)
    )


def read_daily_hydrology_workbook(source: str | Path | BinaryIO) -> pd.DataFrame:
    """Read daily mean inflow from a TV_Ngay workbook."""
    excel = pd.ExcelFile(source)
    records: list[dict[str, object]] = []
    for sheet in excel.sheet_names:
        raw = pd.read_excel(excel, sheet_name=sheet, header=None)
        if raw.empty:
            continue
        plant = next(
            (
                canonical_plant(value)
                for value in raw.iloc[:5, :5].to_numpy().ravel()
                if canonical_plant(value)
            ),
            None,
        )
        if not plant:
            continue
        header_row = next(
            (
                row_idx
                for row_idx in range(min(12, len(raw)))
                if any(_plain(value) == "ngay" for value in raw.iloc[row_idx] if pd.notna(value))
            ),
            None,
        )
        if header_row is None:
            continue
        flow_col = next(
            (
                col_idx
                for col_idx, value in enumerate(raw.iloc[header_row])
                if "luu luong ve ho" in _plain(value)
            ),
            None,
        )
        if flow_col is None:
            continue
        dates = _date_series(raw.iloc[header_row + 1 :, 0])
        flows = pd.to_numeric(raw.iloc[header_row + 1 :, flow_col], errors="coerce")
        valid = dates.notna() & flows.notna()
        for date, flow in zip(dates[valid], flows[valid]):
            records.append(
                {"date": pd.Timestamp(date).normalize(), "plant": plant, "flow": float(flow)}
            )
    if not records:
        return pd.DataFrame(columns=["date", "plant", "flow"])
    return (
        pd.DataFrame(records)
        .groupby(["date", "plant"], as_index=False)["flow"]
        .mean()
        .sort_values(["date", "plant"])
        .reset_index(drop=True)
    )


def load_daily_folders(folder: str | Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Load daily production and hydrology workbooks from their dedicated folders."""
    folder = Path(folder)
    production_frames: list[pd.DataFrame] = []
    hydrology_frames: list[pd.DataFrame] = []
    subfolders = { _plain(path.name): path for path in folder.iterdir() if path.is_dir() }
    production_folder = subfolders.get("san luong")
    hydrology_folder = subfolders.get("thuy van ngay")

    if production_folder:
        for path in sorted(production_folder.glob("*.xls*")):
            try:
                parsed = read_daily_production_workbook(path)
                if not parsed.empty:
                    production_frames.append(parsed)
            except Exception:
                continue
    if hydrology_folder:
        for path in sorted(hydrology_folder.glob("*.xls*")):
            try:
                parsed = read_daily_hydrology_workbook(path)
                if not parsed.empty:
                    hydrology_frames.append(parsed)
            except Exception:
                continue

    production = (
        pd.concat(production_frames, ignore_index=True)
        .groupby(["date", "plant"], as_index=False)["production"]
        .mean()
        if production_frames
        else pd.DataFrame(columns=["date", "plant", "production"])
    )
    hydrology = (
        pd.concat(hydrology_frames, ignore_index=True)
        .groupby(["date", "plant"], as_index=False)["flow"]
        .mean()
        if hydrology_frames
        else pd.DataFrame(columns=["date", "plant", "flow"])
    )
    return production, hydrology


def load_current_hydrology_reports(folder: str | Path) -> pd.DataFrame:
    """Load dated BCTV_LUYKE summaries placed in the dashboard's root folder."""
    frames: list[pd.DataFrame] = []
    for path in sorted(Path(folder).glob("*.xls*")):
        if path.name.startswith("~$"):
            continue
        try:
            parsed = read_cumulative_hydrology_report(path)
        except Exception:
            continue
        if not parsed.empty:
            frames.append(parsed)
    if not frames:
        return pd.DataFrame(columns=["date", "plant", "flow"])
    return (
        pd.concat(frames, ignore_index=True)
        .groupby(["date", "plant"], as_index=False)["flow"]
        .mean()
        .sort_values(["date", "plant"])
        .reset_index(drop=True)
    )


def latest_complete_actual_date(
    daily_production: pd.DataFrame,
    daily_hydrology: pd.DataFrame,
    reported_hydrology: pd.DataFrame | None = None,
) -> pd.Timestamp | None:
    """Return the newest date having production and inflow for all three plants."""
    if daily_production.empty:
        return None
    production_dates = set(
        daily_production.groupby("date")["plant"]
        .nunique()
        .loc[lambda values: values >= len(PLANTS)]
        .index
    )
    hydrology_dates: set[pd.Timestamp] = set()
    if not daily_hydrology.empty:
        hydrology_dates.update(
            daily_hydrology.groupby("date")["plant"]
            .nunique()
            .loc[lambda values: values >= len(PLANTS)]
            .index
        )
    if reported_hydrology is not None and not reported_hydrology.empty:
        hydrology_dates.update(
            reported_hydrology.groupby("date")["plant"]
            .nunique()
            .loc[lambda values: values >= len(PLANTS)]
            .index
        )
    common_dates = production_dates & hydrology_dates
    return pd.Timestamp(max(common_dates)).normalize() if common_dates else None


def current_month_actuals(
    daily_production: pd.DataFrame,
    daily_hydrology: pd.DataFrame,
    reported_hydrology: pd.DataFrame | None,
    actual_as_of: pd.Timestamp | None,
) -> pd.DataFrame:
    """Aggregate production and mean inflow through the selected reporting date."""
    columns = ["year", "month", "plant", "production", "flow"]
    if actual_as_of is None or daily_production.empty:
        return pd.DataFrame(columns=columns)
    actual_as_of = pd.Timestamp(actual_as_of).normalize()
    production = daily_production[
        (daily_production["date"].dt.year == actual_as_of.year)
        & (daily_production["date"].dt.month == actual_as_of.month)
        & (daily_production["date"] <= actual_as_of)
    ]

    report = pd.DataFrame()
    if reported_hydrology is not None and not reported_hydrology.empty:
        report = reported_hydrology[reported_hydrology["date"] == actual_as_of]
    if daily_hydrology.empty:
        daily_flow = pd.DataFrame(columns=["date", "plant", "flow"])
    else:
        daily_flow = daily_hydrology[
            (daily_hydrology["date"].dt.year == actual_as_of.year)
            & (daily_hydrology["date"].dt.month == actual_as_of.month)
            & (daily_hydrology["date"] <= actual_as_of)
        ]

    rows: list[dict[str, object]] = []
    for plant in PLANTS:
        plant_production = production[production["plant"] == plant]["production"]
        report_flow = (
            report[report["plant"] == plant]["flow"]
            if not report.empty
            else pd.Series(dtype=float)
        )
        plant_daily_flow = daily_flow[daily_flow["plant"] == plant]["flow"]
        flow = report_flow.mean() if not report_flow.empty else plant_daily_flow.mean()
        rows.append(
            {
                "year": actual_as_of.year,
                "month": actual_as_of.month,
                "plant": plant,
                "production": plant_production.sum(min_count=1),
                "flow": flow,
            }
        )
    return pd.DataFrame(rows, columns=columns)


def monthly_cumulative_production(
    daily_production: pd.DataFrame,
    year: int,
    month: int,
    through_day: int,
) -> pd.DataFrame:
    """Sum daily production from the first of a month through a calendar day."""
    columns = ["plant", "production", "days"]
    if daily_production.empty:
        return pd.DataFrame(
            {"plant": PLANTS, "production": [np.nan] * len(PLANTS), "days": [0] * len(PLANTS)}
        )[columns]

    last_day = pd.Timestamp(year=year, month=month, day=1).days_in_month
    cutoff = pd.Timestamp(year=year, month=month, day=min(max(int(through_day), 1), last_day))
    monthly = daily_production[
        (daily_production["date"].dt.year == year)
        & (daily_production["date"].dt.month == month)
        & (daily_production["date"] <= cutoff)
    ]
    totals = (
        monthly.groupby("plant", as_index=False)
        .agg(
            production=("production", lambda values: values.sum(min_count=1)),
            days=("date", "nunique"),
        )
    )
    return (
        pd.DataFrame({"plant": PLANTS})
        .merge(totals, on="plant", how="left")
        .assign(days=lambda frame: frame["days"].fillna(0).astype(int))
        [columns]
    )


def _historical_cell_map(worksheet, year: int) -> tuple[int, dict[int, int], dict[int, int]]:
    for header_row in range(1, min(worksheet.max_row, 10) + 1):
        month_positions: dict[int, list[int]] = {month: [] for month in range(1, 13)}
        for column in range(1, worksheet.max_column + 1):
            number = _number(worksheet.cell(header_row, column).value)
            if np.isfinite(number) and number.is_integer() and 1 <= number <= 12:
                month_positions[int(number)].append(column)
        if not all(len(month_positions[month]) >= 2 for month in range(1, 13)):
            continue
        production_columns = {month: month_positions[month][0] for month in range(1, 13)}
        flow_columns = {month: month_positions[month][1] for month in range(1, 13)}
        year_column = production_columns[1] - 1
        for row in range(header_row + 1, worksheet.max_row + 1):
            cell_year = _number(worksheet.cell(row, year_column).value)
            if np.isfinite(cell_year) and int(cell_year) == year:
                return row, production_columns, flow_columns
    raise ValueError(f"Không tìm thấy dòng năm {year} hoặc hai khối tháng 1–12.")


def _set_numeric_cells_in_worksheet(
    worksheet_xml: bytes, values_by_address: dict[str, float]
) -> bytes:
    """Replace selected numeric cells while preserving all other formula caches."""
    spreadsheet_namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", spreadsheet_namespace)
    root = ET.fromstring(worksheet_xml)
    cells = {
        cell.attrib.get("r"): cell
        for cell in root.iter(f"{{{spreadsheet_namespace}}}c")
        if cell.attrib.get("r") in values_by_address
    }
    missing = set(values_by_address) - set(cells)
    if missing:
        raise ValueError(f"Không tìm thấy ô cần cập nhật: {', '.join(sorted(missing))}.")
    for address, value in values_by_address.items():
        cell = cells[address]
        cell.set("t", "n")
        for child_name in ("f", "is"):
            child = cell.find(f"{{{spreadsheet_namespace}}}{child_name}")
            if child is not None:
                cell.remove(child)
        value_node = cell.find(f"{{{spreadsheet_namespace}}}v")
        if value_node is None:
            value_node = ET.SubElement(cell, f"{{{spreadsheet_namespace}}}v")
        value_node.text = format(float(value), ".15g")
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _worksheet_archive_path(workbook_archive: zipfile.ZipFile, sheet_name: str) -> str:
    spreadsheet_namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    office_rel_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    workbook_root = ET.fromstring(workbook_archive.read("xl/workbook.xml"))
    sheet = next(
        (
            item
            for item in workbook_root.iter(f"{{{spreadsheet_namespace}}}sheet")
            if item.attrib.get("name") == sheet_name
        ),
        None,
    )
    if sheet is None:
        raise ValueError(f"Không tìm thấy sheet {sheet_name} trong cấu trúc XLSX.")
    relationship_id = sheet.attrib.get(f"{{{office_rel_namespace}}}id")
    relationships_root = ET.fromstring(
        workbook_archive.read("xl/_rels/workbook.xml.rels")
    )
    relationship = next(
        (
            item
            for item in relationships_root.iter(f"{{{package_rel_namespace}}}Relationship")
            if item.attrib.get("Id") == relationship_id
        ),
        None,
    )
    if relationship is None or not relationship.attrib.get("Target"):
        raise ValueError(f"Không xác định được tệp XML của sheet {sheet_name}.")
    target = relationship.attrib["Target"]
    return target.lstrip("/") if target.startswith("/") else posixpath.normpath(
        posixpath.join("xl", target)
    )


def _patch_xlsx_numeric_cells(
    source: Path,
    destination: Path,
    sheet_name: str,
    values_by_address: dict[str, float],
) -> None:
    """Patch selected XLSX cells without recalculating or clearing other formulas."""
    with zipfile.ZipFile(source, "r") as input_archive:
        worksheet_path = _worksheet_archive_path(input_archive, sheet_name)
        patched_worksheet = _set_numeric_cells_in_worksheet(
            input_archive.read(worksheet_path), values_by_address
        )
        with zipfile.ZipFile(destination, "w") as output_archive:
            for entry in input_archive.infolist():
                content = patched_worksheet if entry.filename == worksheet_path else input_archive.read(entry)
                output_archive.writestr(entry, content)


def update_historical_workbooks(
    folder: str | Path,
    actuals: pd.DataFrame,
    year: int,
    month: int,
) -> tuple[list[dict[str, object]], Path | None]:
    """Write one month's current actuals to all three historical workbooks safely."""
    if not 1 <= month <= 12:
        raise ValueError("Tháng cập nhật phải nằm trong khoảng 1–12.")
    required = actuals.set_index("plant").reindex(PLANTS)
    if required[["production", "flow"]].isna().any().any():
        raise ValueError("Chưa đủ sản lượng và lưu lượng hiện tại của cả 3 nhà máy.")

    prepared: list[dict[str, object]] = []
    folder = Path(folder)
    for plant, filename in HISTORICAL_WORKBOOKS.items():
        path = folder / filename
        if not path.exists():
            raise FileNotFoundError(f"Không tìm thấy {filename}.")
        workbook = load_workbook(path, read_only=True, data_only=False)
        worksheet = next(
            (sheet for sheet in workbook.worksheets if canonical_plant(sheet.title) == plant),
            None,
        )
        if worksheet is None:
            workbook.close()
            raise ValueError(f"Không tìm thấy sheet của {plant} trong {filename}.")
        row, production_columns, flow_columns = _historical_cell_map(worksheet, year)
        production_cell = worksheet.cell(row, production_columns[month])
        flow_cell = worksheet.cell(row, flow_columns[month])
        prepared.append(
            {
                "plant": plant,
                "path": path,
                "sheet_name": worksheet.title,
                "production_address": f"{get_column_letter(production_columns[month])}{row}",
                "flow_address": f"{get_column_letter(flow_columns[month])}{row}",
                "old_production": production_cell.value,
                "old_flow": flow_cell.value,
                "production": float(required.at[plant, "production"]),
                "flow": float(required.at[plant, "flow"]),
            }
        )
        workbook.close()

    changed = [
        item
        for item in prepared
        if not np.isclose(
            _number(item["old_production"]),
            item["production"],
            rtol=1e-12,
            atol=1e-12,
            equal_nan=True,
        )
        or not np.isclose(
            _number(item["old_flow"]),
            item["flow"],
            rtol=1e-12,
            atol=1e-12,
            equal_nan=True,
        )
    ]
    if not changed:
        return [], None

    backup_folder = folder / "backups" / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup_folder.mkdir(parents=True, exist_ok=False)
    temp_paths: list[tuple[Path, Path]] = []
    try:
        for item in prepared:
            path = item["path"]
            shutil.copy2(path, backup_folder / path.name)
        for item in changed:
            path = item["path"]
            temporary = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp.xlsx")
            _patch_xlsx_numeric_cells(
                path,
                temporary,
                item["sheet_name"],
                {
                    item["production_address"]: item["production"],
                    item["flow_address"]: item["flow"],
                },
            )
            temp_paths.append((temporary, path))
        for temporary, path in temp_paths:
            os.replace(temporary, path)
    finally:
        for temporary, _ in temp_paths:
            if temporary.exists():
                temporary.unlink()

    results = [
        {
            "Nhà máy": item["plant"],
            "Sản lượng": item["production"],
            "Lưu lượng": item["flow"],
            "Tệp": item["path"].name,
        }
        for item in changed
    ]
    return results, backup_folder


def read_cumulative_hydrology_report(source: str | Path | BinaryIO) -> pd.DataFrame:
    """Read report date and current-month mean inflow from a BCTV_LUYKE report."""
    excel = pd.ExcelFile(source)
    records: list[dict[str, object]] = []
    for sheet in excel.sheet_names:
        raw = pd.read_excel(excel, sheet_name=sheet, header=None)
        if raw.empty or raw.shape[1] < 4:
            continue

        report_date = None
        for value in raw.iloc[:8, :5].to_numpy().ravel():
            match = re.search(r"(\d{1,2})/(\d{1,2})/((?:19|20)\d{2})", str(value))
            if match:
                report_date = pd.Timestamp(
                    year=int(match.group(3)), month=int(match.group(2)), day=int(match.group(1))
                )
                break
        if report_date is None:
            continue

        current_plant = None
        for row_idx in range(len(raw)):
            label = raw.iloc[row_idx, 1] if raw.shape[1] > 1 else None
            plant = canonical_plant(label)
            if plant:
                current_plant = plant
                continue
            if current_plant and "luu luong tb vao ho" in _plain(label):
                flow = _number(raw.iloc[row_idx, 3])
                if np.isfinite(flow):
                    records.append(
                        {
                            "date": report_date.normalize(),
                            "plant": current_plant,
                            "flow": flow,
                        }
                    )
    if not records:
        return pd.DataFrame(columns=["date", "plant", "flow"])
    return (
        pd.DataFrame(records)
        .groupby(["date", "plant"], as_index=False)["flow"]
        .mean()
        .sort_values(["date", "plant"])
        .reset_index(drop=True)
    )


def _parse_cumulative_hydrology_report(source: str | Path | BinaryIO) -> pd.DataFrame:
    """Read current-month mean inflow from column D of the BCTV_LUYKE report."""
    report = read_cumulative_hydrology_report(source)
    if report.empty:
        return _combine_metric_rows(pd.DataFrame())
    monthly = report.assign(year=report["date"].dt.year, month=report["date"].dt.month)
    return _combine_metric_rows(monthly[["year", "month", "plant", "flow"]])


def _combine_metric_rows(df: pd.DataFrame) -> pd.DataFrame:
    columns = ["year", "month", "plant", "production", "flow"]
    if df.empty:
        return pd.DataFrame(columns=columns)
    for column in columns:
        if column not in df:
            df[column] = np.nan
    df = df.dropna(subset=["year", "month", "plant"])
    df["year"] = df["year"].astype(int)
    df["month"] = df["month"].astype(int)
    return (
        df.groupby(["year", "month", "plant"], as_index=False)[["production", "flow"]]
        .mean()
        .sort_values(["year", "month", "plant"])
        .reset_index(drop=True)
    )


def read_workbook(
    source: str | Path | BinaryIO, cutoff_date: pd.Timestamp | None = None
) -> pd.DataFrame:
    """Read supported Excel layouts without modifying the source workbook."""
    data: bytes | None = None
    if hasattr(source, "read"):
        data = source.read()
    parsers = (
        lambda parser_source: _parse_daily_production_report(parser_source, cutoff_date),
        _parse_cumulative_hydrology_report,
        _parse_long_form,
        _parse_year_sheets,
        _parse_year_rows,
        _parse_dual_matrix,
    )
    frames: list[pd.DataFrame] = []
    for parser in parsers:
        try:
            parser_source = io.BytesIO(data) if data is not None else source
            parsed = parser(parser_source)
            if not parsed.empty:
                frames.append(parsed)
        except Exception:
            continue
    return _combine_metric_rows(pd.concat(frames, ignore_index=True)) if frames else pd.DataFrame()


def load_folder(folder: str | Path) -> tuple[pd.DataFrame, list[str]]:
    folder = Path(folder)
    cutoff_date = None
    cutoff_file = folder / "data_cutoff.txt"
    if cutoff_file.exists():
        try:
            cutoff_date = pd.Timestamp(cutoff_file.read_text(encoding="utf-8").strip())
        except (OSError, ValueError):
            cutoff_date = None
    frames: list[pd.DataFrame] = []
    loaded: list[str] = []
    for path in sorted(folder.glob("*")):
        if path.suffix.lower() not in {".xlsx", ".xls", ".xlsm"} or path.name.startswith("~$"):
            continue
        parsed = read_workbook(path, cutoff_date)
        if not parsed.empty:
            frames.append(parsed)
            loaded.append(path.name)
    return (_combine_metric_rows(pd.concat(frames, ignore_index=True)), loaded) if frames else (pd.DataFrame(), loaded)


def current_period(df: pd.DataFrame, today: pd.Timestamp | None = None) -> tuple[int, int, bool]:
    today = today or pd.Timestamp.today()
    target = df[(df["year"] == today.year) & (df["month"] == today.month)]
    if not target.empty:
        return today.year, today.month, True
    latest = df[["year", "month"]].drop_duplicates().sort_values(["year", "month"]).iloc[-1]
    return int(latest["year"]), int(latest["month"]), False


def prepare_analog_comparison_data(
    df: pd.DataFrame,
    current_year: int,
    cutoff_month: int,
    actual_as_of: pd.Timestamp | None,
    daily_production: pd.DataFrame | None = None,
    daily_hydrology: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """Use like-for-like month-to-date values when the cutoff month is partial.

    The current year keeps its observed cumulative production and mean inflow.
    For historical years, the cutoff month is rebuilt from daily data through
    the same calendar day. Years without complete daily production and inflow
    remain eligible for earlier months but cannot rank through the partial one.
    """
    comparison = df.copy()
    if actual_as_of is None:
        return comparison
    actual_as_of = pd.Timestamp(actual_as_of).normalize()
    if (
        actual_as_of.year != current_year
        or actual_as_of.month != cutoff_month
        or actual_as_of.day >= actual_as_of.days_in_month
    ):
        return comparison

    historical_cutoff = (comparison["year"] != current_year) & (
        comparison["month"] == cutoff_month
    )
    comparison = comparison.loc[~historical_cutoff].copy()
    if (
        daily_production is None
        or daily_production.empty
        or daily_hydrology is None
        or daily_hydrology.empty
    ):
        return comparison

    production_scope = daily_production[
        (daily_production["date"].dt.year != current_year)
        & (daily_production["date"].dt.month == cutoff_month)
        & (daily_production["date"].dt.day <= actual_as_of.day)
    ].assign(year=lambda frame: frame["date"].dt.year)
    flow_scope = daily_hydrology[
        (daily_hydrology["date"].dt.year != current_year)
        & (daily_hydrology["date"].dt.month == cutoff_month)
        & (daily_hydrology["date"].dt.day <= actual_as_of.day)
    ].assign(year=lambda frame: frame["date"].dt.year)
    production_partial = production_scope.groupby(
        ["year", "plant"], as_index=False
    ).agg(
        production=("production", lambda values: values.sum(min_count=1)),
        production_days=("date", "nunique"),
    )
    flow_partial = flow_scope.groupby(["year", "plant"], as_index=False).agg(
        flow=("flow", "mean"),
        flow_days=("date", "nunique"),
    )
    historical_partial = production_partial.merge(
        flow_partial, on=["year", "plant"], how="inner"
    )
    historical_partial = historical_partial[
        (historical_partial["production_days"] >= actual_as_of.day)
        & (historical_partial["flow_days"] >= actual_as_of.day)
    ].assign(month=cutoff_month)
    if historical_partial.empty:
        return comparison

    return (
        pd.concat(
            [
                comparison,
                historical_partial[["year", "month", "plant", "production", "flow"]],
            ],
            ignore_index=True,
        )
        .sort_values(["year", "month", "plant"])
        .reset_index(drop=True)
    )


def rank_analog_years(
    df: pd.DataFrame, year: int, month: int, plant: str | None = None
) -> pd.DataFrame:
    """Rank years using production and flow patterns from January to cutoff month.

    For each plant/metric series, the mean absolute error is normalized by the
    current year's mean absolute value. Candidate years must cover every value
    available in the current January-to-cutoff period.
    """
    scope = df[df["month"].between(1, month)].copy()
    if plant:
        scope = scope[scope["plant"] == plant]
    current = scope[scope["year"] == year].set_index(["plant", "month"])
    if current.empty:
        return pd.DataFrame()

    rows: list[dict[str, float | int]] = []
    for candidate_year in sorted(scope.loc[scope["year"] != year, "year"].unique()):
        candidate = scope[scope["year"] == candidate_year].set_index(["plant", "month"])
        series_errors: list[float] = []
        comparisons = 0
        complete = True
        for plant_name in current.index.get_level_values("plant").unique():
            for metric in ("production", "flow"):
                current_series = current.xs(plant_name, level="plant")[metric].dropna().sort_index()
                if current_series.empty:
                    continue
                candidate_values: list[float] = []
                for series_month in current_series.index:
                    key = (plant_name, series_month)
                    if key not in candidate.index or pd.isna(candidate.at[key, metric]):
                        complete = False
                        break
                    candidate_values.append(float(candidate.at[key, metric]))
                if not complete:
                    break
                candidate_series = np.asarray(candidate_values, dtype=float)
                current_values = current_series.to_numpy(dtype=float)
                scale = max(float(np.mean(np.abs(current_values))), 1e-9)
                series_errors.append(float(np.mean(np.abs(candidate_series - current_values))) / scale)
                comparisons += len(current_values)
            if not complete:
                break
        if complete and series_errors:
            score = float(np.mean(series_errors))
            rows.append(
                {
                    "year": int(candidate_year),
                    "score": score,
                    "similarity": max(0.0, 1.0 - score),
                    "comparisons": comparisons,
                }
            )
    return pd.DataFrame(rows).sort_values(["score", "year"]).reset_index(drop=True) if rows else pd.DataFrame()


def comparison_table(
    df: pd.DataFrame, current_year: int, month: int, analog_year: int, plant: str | None = None
) -> pd.DataFrame:
    scope = df[df["month"].between(1, month)]
    if plant:
        scope = scope[scope["plant"] == plant]
    current = scope[scope["year"] == current_year].set_index(["plant", "month"])
    analog = scope[scope["year"] == analog_year].set_index(["plant", "month"])
    rows = []
    for plant_name in PLANTS:
        for series_month in range(1, month + 1):
            key = (plant_name, series_month)
            if key not in current.index and key not in analog.index:
                continue
            row: dict[str, object] = {"Nhà máy": plant_name, "Tháng": series_month}
            for metric, label in (("production", "Sản lượng"), ("flow", "Lưu lượng")):
                now = current.at[key, metric] if key in current.index else np.nan
                old = analog.at[key, metric] if key in analog.index else np.nan
                row[f"{label} {current_year}"] = now
                row[f"{label} {analog_year}"] = old
                row[f"Chênh lệch {label.lower()} (%)"] = (
                    (float(old) - float(now)) / abs(float(now)) * 100
                    if pd.notna(now) and pd.notna(old) and float(now) != 0
                    else np.nan
                )
            rows.append(row)
    return pd.DataFrame(rows)


def analog_similarity_series(
    df: pd.DataFrame,
    current_year: int,
    cutoff_month: int,
    plant: str,
    analog_year: int,
) -> pd.DataFrame:
    """Build a normalized monthly index for visually comparing two years.

    Production and flow are each divided by the current year's mean for the
    selected period, then averaged. An index of 100 therefore represents the
    current year's average combined production/flow level.
    """
    scope = df[
        (df["plant"] == plant)
        & (df["year"].isin([current_year, analog_year]))
        & (df["month"].between(1, cutoff_month))
    ]
    current = scope[scope["year"] == current_year].set_index("month")
    if current.empty:
        return pd.DataFrame(columns=["Tháng", "Năm", "Chỉ số tương đồng"])

    scales = {
        metric: max(float(current[metric].abs().mean()), 1e-9)
        for metric in ("production", "flow")
        if current[metric].notna().any()
    }
    rows: list[dict[str, object]] = []
    for comparison_year in (current_year, analog_year):
        yearly = scope[scope["year"] == comparison_year].set_index("month")
        for series_month in range(1, cutoff_month + 1):
            if series_month not in yearly.index:
                continue
            normalized = [
                float(yearly.at[series_month, metric]) / scale * 100
                for metric, scale in scales.items()
                if pd.notna(yearly.at[series_month, metric])
            ]
            if normalized:
                rows.append(
                    {
                        "Tháng": series_month,
                        "Năm": str(comparison_year),
                        "Chỉ số tương đồng": float(np.mean(normalized)),
                    }
                )
    return pd.DataFrame(rows)


def annual_projection_table(
    df: pd.DataFrame,
    current_year: int,
    cutoff_month: int,
    actual_as_of: pd.Timestamp | None = None,
    daily_production: pd.DataFrame | None = None,
    ranking_data: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """Combine actual production with analog-year values for the remaining months.

    Months 1 through ``cutoff_month`` come from ``current_year``. Later months
    come from each plant's closest analog year, calculated independently from
    production and flow data available up to the cutoff month.
    """
    columns = ["Nhà máy", "Năm quy luật", *[f"T{month}" for month in range(1, 13)], "Tổng năm"]
    rows: list[dict[str, object]] = []

    actual_as_of = pd.Timestamp(actual_as_of) if actual_as_of is not None else None
    is_partial_month = bool(
        actual_as_of is not None
        and actual_as_of.year == current_year
        and actual_as_of.month == cutoff_month
        and actual_as_of.day < actual_as_of.days_in_month
    )
    ranking_month = cutoff_month - 1 if is_partial_month and cutoff_month > 1 else cutoff_month
    daily_remainder = daily_remaining_production_forecast(
        df,
        daily_production,
        current_year,
        cutoff_month,
        actual_as_of,
        ranking_data,
    )
    ranking_source = ranking_data if ranking_data is not None else df

    for plant in PLANTS:
        ranking = rank_analog_years(
            ranking_source,
            current_year,
            cutoff_month if ranking_data is not None else ranking_month,
            plant,
        )
        analog_year = int(ranking.iloc[0]["year"]) if not ranking.empty else None
        current = (
            df[(df["year"] == current_year) & (df["plant"] == plant)]
            .set_index("month")["production"]
        )
        analog = (
            df[(df["year"] == analog_year) & (df["plant"] == plant)]
            .set_index("month")["production"]
            if analog_year is not None
            else pd.Series(dtype=float)
        )

        row: dict[str, object] = {"Nhà máy": plant, "Năm quy luật": analog_year}
        monthly_values: list[float] = []
        for month in range(1, 13):
            if month < cutoff_month or (month == cutoff_month and not is_partial_month):
                value = current.get(month, np.nan)
            elif month == cutoff_month and is_partial_month:
                actual_value = current.get(month, np.nan)
                analog_value = analog.get(month, np.nan)
                remaining_ratio = (actual_as_of.days_in_month - actual_as_of.day) / actual_as_of.days_in_month
                plant_daily = daily_remainder[daily_remainder["Nhà máy"] == plant]
                daily_value = (
                    float(plant_daily["Sản lượng dự kiến"].sum())
                    if len(plant_daily) == actual_as_of.days_in_month - actual_as_of.day
                    else np.nan
                )
                value = (
                    float(actual_value) + daily_value
                    if pd.notna(actual_value) and np.isfinite(daily_value)
                    else float(actual_value) + float(analog_value) * remaining_ratio
                    if pd.notna(actual_value) and pd.notna(analog_value)
                    else np.nan
                )
            else:
                value = analog.get(month, np.nan)
            numeric_value = float(value) if pd.notna(value) else np.nan
            row[f"T{month}"] = numeric_value
            monthly_values.append(numeric_value)
        row["Tổng năm"] = float(np.nansum(monthly_values)) if any(
            np.isfinite(value) for value in monthly_values
        ) else np.nan
        rows.append(row)

    result = pd.DataFrame(rows, columns=columns)
    total_row: dict[str, object] = {"Nhà máy": "Tổng cộng", "Năm quy luật": None}
    numeric_columns = [f"T{month}" for month in range(1, 13)]
    for column in numeric_columns:
        total_row[column] = result[column].sum(min_count=1)
    total_row["Tổng năm"] = result["Tổng năm"].sum(min_count=1)
    return pd.concat([result, pd.DataFrame([total_row], columns=columns)], ignore_index=True)


def daily_remaining_production_forecast(
    df: pd.DataFrame,
    daily_production: pd.DataFrame | None,
    current_year: int,
    cutoff_month: int,
    actual_as_of: pd.Timestamp | None,
    ranking_data: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """Return analog-year production for each remaining day of a partial month."""
    columns = ["Ngày dự kiến", "Nhà máy", "Năm quy luật", "Sản lượng dự kiến"]
    if daily_production is None or daily_production.empty or actual_as_of is None:
        return pd.DataFrame(columns=columns)
    actual_as_of = pd.Timestamp(actual_as_of)
    if (
        actual_as_of.year != current_year
        or actual_as_of.month != cutoff_month
        or actual_as_of.day >= actual_as_of.days_in_month
    ):
        return pd.DataFrame(columns=columns)

    ranking_source = ranking_data if ranking_data is not None else df
    ranking_month = (
        cutoff_month
        if ranking_data is not None
        else cutoff_month - 1 if cutoff_month > 1 else cutoff_month
    )
    rows: list[dict[str, object]] = []
    for plant in PLANTS:
        ranking = rank_analog_years(ranking_source, current_year, ranking_month, plant)
        if ranking.empty:
            continue
        analog_year = int(ranking.iloc[0]["year"])
        analog_days = daily_production[
            (daily_production["plant"] == plant)
            & (daily_production["date"].dt.year == analog_year)
            & (daily_production["date"].dt.month == cutoff_month)
            & (daily_production["date"].dt.day > actual_as_of.day)
        ]
        for _, daily_row in analog_days.iterrows():
            day = int(pd.Timestamp(daily_row["date"]).day)
            rows.append(
                {
                    "Ngày dự kiến": pd.Timestamp(
                        year=current_year, month=cutoff_month, day=day
                    ),
                    "Nhà máy": plant,
                    "Năm quy luật": analog_year,
                    "Sản lượng dự kiến": float(daily_row["production"]),
                }
            )
    return pd.DataFrame(rows, columns=columns).sort_values(
        ["Ngày dự kiến", "Nhà máy"]
    ).reset_index(drop=True) if rows else pd.DataFrame(columns=columns)
